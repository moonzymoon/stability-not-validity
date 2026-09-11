# -*- coding: utf-8 -*-
"""读取台账结构 (不输出凭据列内容)."""
import openpyxl
import os

BASE = r'D:\0科研\工作1\已投稿'
for f in ('已投稿论文汇总表.xlsx', '已投稿论文汇总表_更新中.xlsx'):
    p = os.path.join(BASE, f)
    try:
        wb = openpyxl.load_workbook(p, read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(min_row=1, max_row=15, values_only=True))
        print('===', f, '| sheet:', ws.title, '| dims:', ws.max_row, 'x',
              ws.max_column)
        hdr = list(rows[0])
        print('列索引-列名:')
        for i, h in enumerate(hdr):
            print(' ', i, h)
        # 输出每行的关键识别列(篇号/标题/期刊/状态), 跳过可能是凭据的列
        cred = {i for i, h in enumerate(hdr)
                if h and any(k in str(h) for k in ('密码', '账号', 'password',
                                                   'Password', '登录'))}
        safe = [i for i in range(len(hdr)) if i not in cred]
        for r_i, row in enumerate(rows[1:], start=2):
            vals = [(hdr[i], row[i]) for i in safe if i < len(row)]
            print(f'row{r_i}:', vals)
        break
    except PermissionError:
        print('===', f, ': LOCKED')
    except FileNotFoundError:
        print('===', f, ': 不存在')
